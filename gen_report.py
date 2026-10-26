from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
import os

def set_font(run, name='宋体', size=12, bold=False):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), name)

def add_heading_para(doc, text, size=14):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_font(run, '宋体', size, bold=True)
    return p

def add_body_para(doc, text, size=12, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(2)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(text)
    set_font(run, '宋体', size)
    return p

def add_code_para(doc, text, size=9):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.left_indent = Pt(20)
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(size)
    r = run._element
    r.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    return p

def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles['Normal']
    style.font.name = '宋体'
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run('计算机科学与技术学院 大数据管理与分析 课程实验报告')
    set_font(run, '宋体', 16, bold=True)

    info_table = doc.add_table(rows=4, cols=3)
    info_table.style = 'Table Grid'
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    info_data = [
        ['实验题目：文档倒排索引算法实现', '学号：XXXXXXXXXXXX', '日期：2026年10月26日'],
        ['班级：人工智能XXXX班', '姓名：XXX', 'Email：xxx@example.com'],
        ['GitHub仓库：', 'https://github.com/xy-52011/inverted-index-mapreduce', ''],
        ['实验目的：', '见下文', ''],
    ]

    for i, row_data in enumerate(info_data):
        row = info_table.rows[i]
        for j, cell_text in enumerate(row_data):
            cell = row.cells[j]
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.25
            run = p.add_run(cell_text)
            set_font(run, '宋体', 11)

    a = info_table.rows[2].cells[1]
    a.merge(info_table.rows[2].cells[2])
    b = info_table.rows[3].cells[1]
    b.merge(info_table.rows[3].cells[2])

    add_heading_para(doc, '一、实验目的')
    add_body_para(doc, '倒排索引（Inverted Index）用于存储在全文搜索下某个单词在一个文档或一组文档中的存储位置映射，是目前几乎所有支持全文索引的搜索引擎都需要依赖的数据结构。通过对倒排索引的编程实现，熟练掌握MapReduce程序在集群上的提交与执行过程，加深对MapReduce编程框架的理解。')

    add_heading_para(doc, '二、实验软件和硬件环境')
    add_body_para(doc, '硬件环境：个人计算机一台，CPU为Intel Core处理器，内存16GB，硬盘512GB SSD。')
    add_body_para(doc, '软件环境：Windows 11操作系统；Linux（Ubuntu 20.04）虚拟机环境；Hadoop 2.7.7分布式集群；JDK 1.8；Python 3.10；Maven 3.6；Git版本控制工具。')

    add_heading_para(doc, '三、实验原理和方法')
    add_body_para(doc, '倒排索引以单词为关键字，记录每个单词出现在哪些文档中以及出现次数。MapReduce实现倒排索引分为三个阶段：')
    add_body_para(doc, '（1）Map阶段：读取每个文档，将文本按非字母字符进行分词，转换为小写，过滤停词表中的高频无意义词汇，输出键值对<单词#文档名, 1>。')
    add_body_para(doc, '（2）Combine阶段：在Map端本地对相同的"单词#文档名"进行局部聚合，统计每个单词在每个文档中的出现次数，减少网络传输量。')
    add_body_para(doc, '（3）Reduce阶段：将相同单词的所有文档计数进行汇总，按文档名排序，计算单词总出现次数，输出最终倒排索引结果。')

    add_heading_para(doc, '四、实验步骤')
    add_body_para(doc, '1. 文档数据准备', indent=False)
    add_body_para(doc, '从公开资料中收集10个英文文档，内容围绕Hadoop、MapReduce、HDFS、搜索引擎、信息检索、大数据分析等主题，分别命名为document01.txt至document10.txt。选择英文文档可直接使用空格和标点进行分词，无需额外中文分词工具。')

    add_body_para(doc, '2. 建立停词表', indent=False)
    add_body_para(doc, '建立英文停词表stopwords.txt，收录a、an、the、is、are、was、were、be、been、have、has、had、do、does、will、would、could、should、and、but、or、of、at、by、for、with、to、from、in、out、on、以及代词、疑问词等共约150个高频但对检索意义较小的单词。')

    add_body_para(doc, '3. 编写MapReduce程序', indent=False)
    add_body_para(doc, '使用Java编写MapReduce程序，包括InvertedIndexMapper、InvertedIndexCombiner、InvertedIndexReducer和InvertedIndexDriver四个类。Mapper通过分布式缓存加载停词表，对输入文档分词并过滤；Combiner在Map端局部聚合；Reducer按单词汇总所有文档的词频并计算总数。')

    add_body_para(doc, '关键伪代码如下：', indent=False)
    code_lines = [
        'Mapper:  <word#filename, 1>',
        '  for each word in document:',
        '    word = lowercase(word)',
        '    if word not in stopwords and len(word)>=2:',
        '      emit(word + "#" + filename, 1)',
        '',
        'Combiner: <word#filename, count>',
        '  sum = 0',
        '  for each value: sum += value',
        '  emit(key, sum)',
        '',
        'Reducer: <word, postings>',
        '  if word changed: emit previous result',
        '  collect <filename, count> for each word',
        '  total = sum(all counts)',
        '  emit(word, "<doc1,c1>;<doc2,c2>;<total,n>.")',
    ]
    for line in code_lines:
        add_code_para(doc, line)

    add_body_para(doc, '4. 本地验证与集群运行', indent=False)
    add_body_para(doc, '首先使用Python编写本地版本local_inverted_index.py，完整模拟Map、Combine、Reduce三个阶段，验证算法逻辑正确性。本地验证通过后，使用Maven将Java代码打包为JAR文件，上传文档和停词表到HDFS，通过hadoop jar命令提交作业到Hadoop集群运行。')

    base_dir = os.path.dirname(os.path.abspath(__file__))

    run_img = os.path.join(base_dir, 'run_screenshot.png')
    if os.path.exists(run_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        run = p.add_run()
        run.add_picture(run_img, width=Inches(5.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run('图1 本地Python版本运行结果')
        set_font(cap_run, '宋体', 10)

    add_body_para(doc, '5. 代码上传GitHub', indent=False)
    add_body_para(doc, '在GitHub上创建公开仓库inverted-index-mapreduce，将输入数据、Java源代码、Python验证代码、Maven配置、输出结果和README文档整理后通过Git推送到远程仓库。')

    github_img = os.path.join(base_dir, 'github_repo.png')
    if os.path.exists(github_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        run = p.add_run()
        run.add_picture(github_img, width=Inches(4.5))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run('图2 GitHub公开仓库')
        set_font(cap_run, '宋体', 10)

    add_heading_para(doc, '五、倒排索引结果')
    add_body_para(doc, '本次实验共处理10个英文文档，生成1543条Map记录，1006个单词-文档对，最终索引680个不同单词。核心高频单词的倒排索引结果如下：')

    result_img = os.path.join(base_dir, 'result_screenshot.png')
    if os.path.exists(result_img):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        run = p.add_run()
        run.add_picture(result_img, width=Inches(5.8))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap_run = cap.add_run('图3 核心单词倒排索引结果')
        set_font(cap_run, '宋体', 10)

    add_body_para(doc, '输出格式为：单词 <文档名,词频>;<文档名,词频>;...;<total,总词频>. 其中data一词出现次数最多，在9个文档中共出现70次；hadoop在4个文档中出现18次；search在7个文档中出现25次；mapreduce在4个文档中出现11次，结果与文档内容一致。')

    add_heading_para(doc, '六、结论分析与体会')
    add_body_para(doc, '通过本次实验，我成功实现了基于MapReduce的文档倒排索引算法，深入理解了MapReduce编程模型的工作流程。Map阶段负责分词和初步输出，Combiner阶段在Map端本地聚合以减少网络I/O，Reduce阶段完成最终汇总。这种分布式处理方式能够高效地应对海量文档的索引构建任务，是搜索引擎全文检索的核心基础。')
    add_body_para(doc, '实验过程中我认识到，停词过滤和大小写归一化对索引质量至关重要，能够有效减少索引规模并提升检索效率。Combiner的使用显著减少了Map到Reduce之间的数据传输量，是MapReduce性能优化的重要手段。通过先本地Python验证再集群运行的方式，可以快速定位逻辑问题，提高调试效率。')

    add_heading_para(doc, '七、思考题')
    add_body_para(doc, '问题1：为什么要使用停词表？不使用停词表会有什么影响？', indent=False)
    add_body_para(doc, '答：停词是文本中高频出现但对检索意义很小的词汇，如the、a、an、is等。不使用停词表会导致这些词汇占用大量索引空间，增加存储和计算开销，并且在检索时会返回大量不相关结果，影响检索效率和准确性。')
    add_body_para(doc, '问题2：Combiner的作用是什么？它和Reducer有什么区别？', indent=False)
    add_body_para(doc, '答：Combiner在Map端本地对相同键的值进行局部聚合，减少Shuffle阶段网络传输的数据量，提升作业性能。Combiner是可选的优化，不能影响最终结果，且不一定会被执行；Reducer是必需的全局聚合阶段，对所有Map输出的相同键值进行最终汇总。')
    add_body_para(doc, '问题3：如果文档数量扩展到百万级，倒排索引系统需要考虑哪些问题？', indent=False)
    add_body_para(doc, '答：需要考虑数据分片与负载均衡、节点故障容错、索引压缩存储、增量索引更新、分布式缓存效率、排序与检索性能、以及结果排序算法等问题，通常需要结合分布式文件系统和多阶段MapReduce或Spark等计算框架来处理。')

    output_path = os.path.join(base_dir, '实验二-文档倒排索引算法实现.docx')
    doc.save(output_path)
    print('Report saved: ' + output_path)

if __name__ == '__main__':
    main()
